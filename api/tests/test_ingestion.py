from io import BytesIO

import pytest
from openpyxl import Workbook

from app.core.errors import DomainError
from app.ingestion.parser import parse_upload, validate_mapped_rows


def test_csv_bom_quotes_identifier_stripping_and_normalization():
    raw = '\ufefftask_description,hours_per_week,tools_used,skills_used,email\n"Help, listen, and respond",2,"calendar; CRM","care; writing",private@example.test\n'.encode()
    parsed = parse_upload("tasks.csv", raw)
    sheet = parsed["sheets"]["CSV"]
    assert sheet["headers"] == ["task_description", "hours_per_week", "tools_used", "skills_used"]
    assert "email" in sheet["notices"][0]
    mapping = {"task_description": 0, "hours_per_week": 1, "tools_used": 2, "skills_used": 3}
    accepted, errors = validate_mapped_rows(sheet["headers"], sheet["rows"], mapping)
    assert not errors
    assert accepted[0]["task_description"] == "Help, listen, and respond"
    assert accepted[0]["tools_used"] == ["calendar", "CRM"]


def test_csv_bad_size_invalid_encoding_malformed_quotes_and_row_limit():
    with pytest.raises(DomainError) as too_large:
        parse_upload("a.csv", b"x" * (5 * 1024 * 1024 + 1))
    assert too_large.value.status_code == 413
    with pytest.raises(DomainError):
        parse_upload("a.csv", b"\xff")
    with pytest.raises(DomainError):
        parse_upload("a.csv", b'task_description,hours_per_week\n"unfinished,2')
    content = "task_description,hours_per_week\n" + "Task,1\n" * 61
    with pytest.raises(DomainError) as rows:
        parse_upload("a.csv", content.encode())
    assert rows.value.status_code == 422


def test_mapping_errors_keep_source_row_numbers():
    headers = ["task_description", "hours_per_week", "frequency", "involves_people"]
    rows = [["Good task", "3", "weekly", "no"], ["Bad hours", "not a number", "yearly", "maybe"]]
    accepted, errors = validate_mapped_rows(
        headers,
        rows,
        {"task_description": 0, "hours_per_week": 1, "frequency": 2, "involves_people": 3},
    )
    assert len(accepted) == 1 and accepted[0]["involves_people"] is False
    assert errors[0]["row"] == 3


def test_xlsx_sheet_selection_formula_policy_and_corruption():
    workbook = Workbook()
    first = workbook.active
    first.title = "Week A"
    first.append(["task_description", "hours_per_week"])
    first.append(["Write report", 2])
    second = workbook.create_sheet("Week B")
    second.append(["task_description", "hours_per_week"])
    second.append(["Review report", 1])
    buffer = BytesIO()
    workbook.save(buffer)
    parsed = parse_upload("work.xlsx", buffer.getvalue())
    assert parsed["sheet_names"] == ["Week A", "Week B"]
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["task_description", "hours_per_week"])
    sheet.append(["Formula", None])
    sheet["B2"] = "=1+1"
    buffer = BytesIO()
    workbook.save(buffer)
    with pytest.raises(DomainError, match="Formula cells"):
        parse_upload("formula.xlsx", buffer.getvalue())
    with pytest.raises(DomainError, match="could not be read"):
        parse_upload("corrupt.xlsx", b"not a zip archive")
