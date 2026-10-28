import csv
import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from app.core.config import get_settings
from app.core.errors import DomainError

ALIASES = {
    "task_description": ["task description", "task", "task name", "description", "activity"],
    "hours_per_week": ["hours per week", "weekly hours", "hours", "time per week"],
    "frequency": ["frequency"],
    "tools_used": ["tools used", "tools"],
    "skills_used": ["skills used", "skills"],
    "involves_people": ["involves people", "people interaction"],
    "physical_presence_required": ["physical presence required", "in person required"],
    "decision_authority": ["decision authority"],
    "notes": ["notes"],
}
IDENTIFIER_HEADERS = {
    "name",
    "full name",
    "first name",
    "last name",
    "email",
    "email address",
    "phone",
    "phone number",
    "ssn",
    "social security number",
    "address",
    "employer",
    "company",
}


def suggest_mappings(headers: list[str]) -> tuple[dict, list[str]]:
    mapping, notices = {}, []
    for field, aliases in ALIASES.items():
        matches = [
            i
            for i, header in enumerate(headers)
            if header.strip().lower().replace("_", " ") in aliases
        ]
        if len(matches) == 1:
            mapping[field] = matches[0]
        elif len(matches) > 1:
            notices.append(f"More than one column could map to {field}; choose one manually.")
    if "task_description" not in mapping or "hours_per_week" not in mapping:
        notices.append("Map task description and weekly hours before confirming.")
    return mapping, notices


def _clean_sheet(headers, rows):
    kept = []
    removed = []
    for index, header in enumerate(headers):
        normalized = str(header or "").strip().lower().replace("_", " ")
        if normalized in IDENTIFIER_HEADERS:
            removed.append(str(header))
            continue
        kept.append(index)
    clean_headers = [str(headers[i] or "").strip() for i in kept]
    clean_rows = [
        [str(row[i] if i < len(row) and row[i] is not None else "").strip() for i in kept]
        for row in rows
    ]
    notices = []
    if removed:
        notices.append(
            "Identifier-looking columns were removed: "
            + ", ".join(removed)
            + ". Free-text descriptions may still identify people."
        )
    if any(not header for header in clean_headers):
        notices.append("One or more columns have a missing header.")
    if len({h.casefold() for h in clean_headers}) != len(clean_headers):
        notices.append("Conflicting duplicate headers are present; choose mappings carefully.")
    return clean_headers, clean_rows, notices


def parse_upload(filename: str, content: bytes) -> dict:
    settings = get_settings()
    if len(content) > settings.max_upload_bytes:
        raise DomainError("upload_too_large", "Files must be 5 MB or smaller.", 413)
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = content.decode("utf-8-sig")
            data = []
            scanned = 0
            cell_count = 0
            for row in csv.reader(io.StringIO(text, newline=""), strict=True):
                scanned += 1
                cell_count += len(row)
                if scanned > 1000 or cell_count > settings.max_workbook_cells:
                    raise DomainError(
                        "csv_too_complex",
                        "The CSV contains too many rows or cells to process safely.",
                        413,
                    )
                if len(row) > 200:
                    raise DomainError(
                        "too_many_columns", "A sheet cannot have more than 200 columns.", 422
                    )
                data.append(row)
        except DomainError:
            raise
        except (UnicodeDecodeError, csv.Error) as exc:
            raise DomainError("invalid_csv", f"Could not parse UTF-8 CSV: {exc}", 422) from exc
        if not data:
            raise DomainError("empty_file", "The CSV file has no content.", 422)
        sheets = {"CSV": data}
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                expanded = sum(item.file_size for item in archive.infolist())
                if expanded > settings.max_workbook_expanded_bytes:
                    raise DomainError(
                        "workbook_too_large",
                        "The expanded workbook exceeds the processing limit.",
                        413,
                    )
                if any(
                    item.filename.lower().endswith("vbaproject.bin") for item in archive.infolist()
                ):
                    raise DomainError(
                        "macros_unsupported", "Macro-enabled content is not accepted.", 422
                    )
            workbook = load_workbook(
                io.BytesIO(content), read_only=True, data_only=False, keep_links=False
            )
            sheets = {}
            cell_count = 0
            for sheet in workbook.worksheets:
                values = []
                for row in sheet.iter_rows():
                    cell_count += len(row)
                    if cell_count > settings.max_workbook_cells:
                        raise DomainError(
                            "workbook_too_complex",
                            "The workbook contains too many cells to process safely.",
                            413,
                        )
                    if any(cell.data_type == "f" for cell in row):
                        raise DomainError(
                            "formula_unsupported",
                            "Formula cells are not accepted. Replace formulas with fixed values; formulas and macros are never evaluated.",
                            422,
                        )
                    values.append([cell.value for cell in row])
                sheets[sheet.title] = values
        except DomainError:
            raise
        except Exception as exc:
            raise DomainError("invalid_xlsx", "The workbook could not be read.", 422) from exc
    else:
        raise DomainError("unsupported_file", "Choose a CSV or XLSX workbook.", 422)
    result = {}
    for name, matrix in sheets.items():
        while matrix and not any(str(value or "").strip() for value in matrix[-1]):
            matrix.pop()
        if not matrix:
            result[name] = {
                "headers": [],
                "rows": [],
                "notices": ["This worksheet is empty."],
                "errors": [],
            }
            continue
        headers, rows, notices = _clean_sheet(matrix[0], matrix[1:])
        if len(rows) > 60:
            raise DomainError(
                "row_limit",
                "Each worksheet may contain at most 60 task rows.",
                422,
                {"sheet": name, "rows": len(rows)},
            )
        if len(rows) + 1 > settings.max_workbook_cells:
            raise DomainError("workbook_too_complex", "The worksheet has too many cells.", 413)
        if not headers:
            raise DomainError("missing_headers", "The file must include a header row.", 422)
        if len({header.casefold() for header in headers}) != len(headers):
            notices.append("Duplicate headers must be resolved manually before import.")
        mapping, mapping_notices = suggest_mappings(headers)
        result[name] = {
            "headers": headers,
            "rows": rows,
            "mapping": mapping,
            "notices": notices + mapping_notices,
            "errors": [],
        }
    return {"sheets": result, "sheet_names": list(result)}


def validate_mapped_rows(
    headers: list[str], rows: list[list[str]], mapping: dict
) -> tuple[list[dict], list[dict]]:
    from app.tasks.schemas import TaskInput

    accepted, errors = [], []
    for row_number, row in enumerate(rows, start=2):
        if not any(str(value).strip() for value in row):
            continue

        def value(field):
            index = mapping.get(field)
            return "" if index is None else (str(row[index]).strip() if index < len(row) else "")

        raw = {
            "task_description": value("task_description"),
            "hours_per_week": value("hours_per_week"),
            "frequency": value("frequency") or None,
            "tools_used": [v.strip() for v in value("tools_used").split(";") if v.strip()] or None,
            "skills_used": [v.strip() for v in value("skills_used").split(";") if v.strip()]
            or None,
            "notes": value("notes") or None,
        }
        for field in ("involves_people", "physical_presence_required"):
            text = value(field).casefold()
            raw[field] = True if text == "yes" else False if text == "no" else None
        decision = value("decision_authority")
        if decision:
            try:
                raw["decision_authority"] = int(decision)
            except ValueError:
                raw["decision_authority"] = decision
        try:
            accepted.append(TaskInput.model_validate(raw).model_dump(mode="json"))
        except Exception as exc:
            errors.append(
                {"row": row_number, "message": "; ".join(error["msg"] for error in exc.errors())}
            )
    if len(accepted) > 60:
        raise DomainError(
            "row_limit", "A confirmed import can contain at most 60 valid tasks.", 422
        )
    return accepted, errors
